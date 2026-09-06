#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
tests/test_build_metre.py — Tests unitaires pour build_metre.py

Valide que le classeur Excel est bien généré avec les 2 feuilles :
  1. "Detail quontitafif fondation" — volumes beton
  2. "Armatures" — ferraillage avec ventilation par diametre
"""
import json
import os
import sys
import unittest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from openpyxl import load_workbook

from build_metre import MetreGenerator, _detect_diameters

PLAN_PATH = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
                         "sample_plan_data.json")
OUTPUT_DIR = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
                          "output")
OUTPUT_FILE = os.path.join(OUTPUT_DIR, "test_build_metre.xlsx")


class TestDetectDiameters(unittest.TestCase):
    """Test la détection dynamique des diamètres."""

    def setUp(self):
        with open(PLAN_PATH, encoding="utf-8") as f:
            self.plan = json.load(f)

    def test_diameters_detected(self):
        diams = _detect_diameters(self.plan)
        self.assertIn(6, diams)
        self.assertIn(12, diams)
        self.assertIn(14, diams)
        self.assertIn(16, diams)

    def test_diameters_sorted(self):
        diams = _detect_diameters(self.plan)
        self.assertEqual(diams, sorted(diams))

    def test_diameters_from_plan(self):
        diams = _detect_diameters(self.plan)
        self.assertIn(6, diams)
        self.assertIn(8, diams)
        self.assertIn(12, diams)
        self.assertIn(14, diams)
        self.assertIn(16, diams)


class TestMetreGenerator(unittest.TestCase):
    """Test le générateur de métré — Feuille Detail quantitatif."""

    @classmethod
    def setUpClass(cls):
        with open(PLAN_PATH, encoding="utf-8") as f:
            cls.plan = json.load(f)
        os.makedirs(OUTPUT_DIR, exist_ok=True)
        gen = MetreGenerator(cls.plan, moteur="historique")
        gen.generer(OUTPUT_FILE)
        cls.wb = load_workbook(OUTPUT_FILE)
        # Feuille 1 = Detail quantitatif
        cls.ws_detail = cls.wb[cls.wb.sheetnames[0]]
        # Feuille 2 = Armatures
        cls.ws_arm = cls.wb[cls.wb.sheetnames[1]]

    def test_workbook_created(self):
        self.assertTrue(os.path.exists(OUTPUT_FILE))

    def test_two_sheets_exist(self):
        self.assertEqual(len(self.wb.sheetnames), 2)
        self.assertIn("Detail quontitafif fondation", self.wb.sheetnames)
        self.assertIn("Armatures", self.wb.sheetnames)

    def test_project_name_in_detail_sheet(self):
        val = self.ws_detail.cell(row=4, column=3).value
        self.assertEqual(val, "Projet Générique")

    def test_detail_headers_row(self):
        headers = [self.ws_detail.cell(row=6, column=c).value for c in range(1, 12)]
        self.assertEqual(headers[0], "N°")
        self.assertIn("Désignation", str(headers[1]))
        self.assertEqual(headers[5], "N")
        self.assertIn("Longueur", str(headers[6]))
        self.assertIn("Largeur", str(headers[7]))
        self.assertIn("Hauteur", str(headers[8]))

    def test_terassement_poste_exists(self):
        found = False
        for row in range(1, self.ws_detail.max_row + 1):
            val = self.ws_detail.cell(row=row, column=1).value
            if val and str(val) == "3":
                found = True
                break
        self.assertTrue(found, "Poste 3 Terrassement non trouvé")

    def test_beton_proprete_poste_exists(self):
        found = False
        for row in range(1, self.ws_detail.max_row + 1):
            val = self.ws_detail.cell(row=row, column=1).value
            if val and str(val) == "15":
                found = True
                break
        self.assertTrue(found, "Poste 15 Béton de propreté non trouvé")

    def test_beton_arme_poste_exists(self):
        found = False
        for row in range(1, self.ws_detail.max_row + 1):
            val = self.ws_detail.cell(row=row, column=1).value
            if val and str(val) == "16":
                found = True
                break
        self.assertTrue(found, "Poste 16 Béton armé non trouvé")

    def test_all_semelle_ids_in_armatures(self):
        """Vérifie que S101..S106 sont tous présents dans Armatures."""
        ids_found = set()
        for row in range(1, self.ws_arm.max_row + 1):
            val = self.ws_arm.cell(row=row, column=1).value
            if val and str(val).startswith("S"):
                ids_found.add(val)
        for sid in ["S101", "S102", "S103", "S104", "S105", "S106"]:
            self.assertIn(sid, ids_found, f"Semelle {sid} introuvable")

    def test_all_poteau_ids_in_armatures(self):
        """Vérifie que P101..P104 sont tous présents."""
        ids_found = set()
        for row in range(1, self.ws_arm.max_row + 1):
            val = self.ws_arm.cell(row=row, column=1).value
            if val and str(val).startswith("P") and not str(val).startswith("PO"):
                ids_found.add(val)
        for pid in ["P101", "P102", "P103", "P104"]:
            self.assertIn(pid, ids_found, f"Poteau {pid} introuvable")

    def test_all_poutre_ids_in_armatures(self):
        """Vérifie que PO1..PO3 sont tous présents."""
        ids_found = set()
        for row in range(1, self.ws_arm.max_row + 1):
            val = self.ws_arm.cell(row=row, column=1).value
            if val and str(val).startswith("PO"):
                ids_found.add(val)
        for pid in ["PO1", "PO2", "PO3"]:
            self.assertIn(pid, ids_found, f"Poutre {pid} introuvable")

    def test_dimensions_from_catalogue(self):
        """Vérifie que les dimensions viennent du catalogue."""
        # S101 est S_TYPE_A : a=1.20, b=1.20, h=0.30
        for row in range(1, self.ws_arm.max_row + 1):
            if self.ws_arm.cell(row=row, column=1).value == "S101":
                self.assertAlmostEqual(self.ws_arm.cell(row=row, column=4).value, 1.20, places=2)
                self.assertAlmostEqual(self.ws_arm.cell(row=row, column=5).value, 1.20, places=2)
                self.assertAlmostEqual(self.ws_arm.cell(row=row, column=6).value, 0.30, places=2)
                break


class TestFormulasArmatures(unittest.TestCase):
    """Test que les formules Excel Armatures sont correctement écrites."""

    @classmethod
    def setUpClass(cls):
        with open(PLAN_PATH, encoding="utf-8") as f:
            cls.plan = json.load(f)
        os.makedirs(OUTPUT_DIR, exist_ok=True)
        gen = MetreGenerator(cls.plan, moteur="historique")
        gen.generer(OUTPUT_FILE)
        cls.wb = load_workbook(OUTPUT_FILE)
        cls.ws_arm = cls.wb["Armatures"]

    def test_armature_headers(self):
        """En-têtes Armatures : A=OUVRAGES, G=N° ELE, H=NOMB B, I=DIAM, J=LONG."""
        h = {c: self.ws_arm.cell(row=1, column=c).value for c in range(1, 11)}
        self.assertIn("O U V R A G E S", str(h[1]))
        self.assertEqual(h[7], "N° ELE")
        self.assertEqual(h[8], "NOMB B")
        self.assertEqual(h[9], "DIAM")
        self.assertEqual(h[10], "LONG")

    def test_acier_tor_headers(self):
        """Colonne K doit avoir T6, L=T8, etc."""
        self.assertIn("T 6", str(self.ws_arm.cell(row=2, column=11).value))
        self.assertIn("T 14", str(self.ws_arm.cell(row=2, column=14).value))

    def test_semelle_longueur_formula_is_formula(self):
        """La longueur d'armature semelle doit être une formule."""
        for row in range(1, self.ws_arm.max_row + 1):
            ouvrage = self.ws_arm.cell(row=row, column=1).value
            val = self.ws_arm.cell(row=row, column=10).value
            if ouvrage and "Armature INF" in str(ouvrage) and val:
                self.assertTrue(str(val).startswith("="),
                                f"Row {row}: LONG devrait être une formule, got {val}")
                break

    def test_ventilation_formula_is_if(self):
        """Les colonnes de ventilation (K+) doivent contenir des formules IF."""
        for row in range(1, self.ws_arm.max_row + 1):
            ouvrage = self.ws_arm.cell(row=row, column=1).value
            if ouvrage and "Armature INF X" in str(ouvrage):
                for col in range(11, self.ws_arm.max_column + 1):
                    val = self.ws_arm.cell(row=row, column=col).value
                    if val:
                        self.assertTrue(str(val).startswith("=IF"),
                                        f"Row {row} Col {col}: ventilation devrait être IF, got {val}")
                break

    def test_cadres_formula_roundup(self):
        """Le nombre de cadres poteaux doit contenir ROUNDUP."""
        found_cadre = False
        for row in range(1, self.ws_arm.max_row + 1):
            val = self.ws_arm.cell(row=row, column=8).value
            if val and isinstance(val, str) and "ROUNDUP" in val:
                found_cadre = True
                self.assertIn("ROUNDUP", val)
                self.assertIn("+2", val)
                break
        self.assertTrue(found_cadre, "Aucune formule ROUNDUP pour cadres trouvée")

    def test_totaux_section_exists(self):
        """Vérifie la présence des lignes de totaux."""
        found_longueur = False
        found_poids_ml = False
        found_poids_partiels = False
        found_poids_total = False
        for row in range(1, self.ws_arm.max_row + 1):
            val = self.ws_arm.cell(row=row, column=1).value
            if val:
                if "LONGUEUR TOTALE" in str(val):
                    found_longueur = True
                elif "POIDS / ML" in str(val):
                    found_poids_ml = True
                elif "POIDS PARTIELS" in str(val):
                    found_poids_partiels = True
                elif "POIDS TOTAL" in str(val):
                    found_poids_total = True
        self.assertTrue(found_longueur, "Ligne LONGUEUR TOTALE introuvable")
        self.assertTrue(found_poids_ml, "Ligne POIDS / ML introuvable")
        self.assertTrue(found_poids_partiels, "Ligne POIDS PARTIELS introuvable")
        self.assertTrue(found_poids_total, "Ligne POIDS TOTAL introuvable")

    def test_longueur_totale_has_sum_formula(self):
        """LONGUEUR TOTALE doit contenir des formules SUM."""
        for row in range(1, self.ws_arm.max_row + 1):
            val = self.ws_arm.cell(row=row, column=1).value
            if val and "LONGUEUR TOTALE" in str(val):
                has_sum = False
                for col in range(11, self.ws_arm.max_column + 1):
                    cell_val = self.ws_arm.cell(row=row, column=col).value
                    if cell_val and "SUM" in str(cell_val):
                        has_sum = True
                        break
                self.assertTrue(has_sum, "LONGUEUR TOTALE sans formule SUM")
                break

    def test_poids_ml_has_formula(self):
        """POIDS / ML doit contenir des formules d²/162."""
        for row in range(1, self.ws_arm.max_row + 1):
            val = self.ws_arm.cell(row=row, column=1).value
            if val and "POIDS / ML" in str(val):
                for col in range(11, self.ws_arm.max_column + 1):
                    cell_val = self.ws_arm.cell(row=row, column=col).value
                    if cell_val:
                        self.assertIn("/162", str(cell_val),
                                        f"POIDS/ML col {col} devrait contenir /162")
                break

    def test_no_hardcoded_project_values(self):
        """Aucune valeur hardcodée du projet MZINDA/Youssoufia."""
        for row in range(1, self.ws_arm.max_row + 1):
            for col in range(1, self.ws_arm.max_column + 1):
                val = self.ws_arm.cell(row=row, column=col).value
                if val and isinstance(val, str):
                    self.assertNotIn("Youssoufia", val)
                    self.assertNotIn("MZINDA", val)


if __name__ == "__main__":
    unittest.main()
