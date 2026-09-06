#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
tests/test_modele_metre.py -- Classeur modele standardise (template vierge).

Verifie : 4 feuilles interconnectees, charte graphique, formules Excel en
majuscules, formats stricts, zero fusion hors cartouche (lignes 1-5),
AUCUNE donnee projet en dur, VBA Module1.bas complet.
"""
import os
import sys
import unittest

import openpyxl

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
VBA_PATH = os.path.join(ROOT, "vba", "Module1.bas")

DIAMS = [6, 8, 10, 12, 14, 16, 20, 25, 32]
POIDS_NOMINAUX = [0.222, 0.395, 0.617, 0.888, 1.208, 1.578,
                  2.466, 3.853, 6.313]


def _generer(tmp_path):
    from generators.modele_metre import generer_modele
    out = os.path.join(str(tmp_path), "modele.xlsx")
    generer_modele(out)
    assert os.path.exists(out)
    return openpyxl.load_workbook(out, data_only=False)


class TestStructure(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        import tempfile
        cls._td = tempfile.TemporaryDirectory()
        cls.addClassCleanup(cls._td.cleanup)
        from generators.modele_metre import generer_modele
        cls.out = os.path.join(cls._td.name, "modele.xlsx")
        generer_modele(cls.out)
        cls.wb = openpyxl.load_workbook(cls.out, data_only=False)

    def test_quatre_feuilles(self):
        self.assertEqual(self.wb.sheetnames,
                         ["01_Detail_Quantitatif", "02_Armatures",
                          "03_Attachement_Ferraillage", "04_GO_Attachement"])

    def test_aucune_fusion_hors_cartouche(self):
        """Fusions reservees aux lignes 1-5 (cartouche)."""
        for ws in self.wb.worksheets:
            for rng in ws.merged_cells.ranges:
                self.assertLessEqual(rng.max_row, 5,
                                     f"{ws.title} : fusion ligne {rng}")

    def test_charte_en_tetes(self):
        ws = self.wb["01_Detail_Quantitatif"]
        c = ws.cell(6, 1)
        self.assertEqual(c.font.name, "Segoe UI")
        self.assertEqual(c.font.size, 11)
        self.assertTrue(c.font.bold)
        self.assertEqual(c.fill.start_color.rgb[-6:], "1B365D")
        c = ws.cell(1, 1)
        self.assertEqual(c.font.size, 13)

    def test_formats_stricts(self):
        ws = self.wb["02_Armatures"]
        self.assertEqual(ws.cell(4, 10).number_format, "#,##0.00")
        self.assertEqual(ws.cell(4, 9).number_format, "0")
        self.assertEqual(ws.cell(4, 8).number_format, "0")
        ws1 = self.wb["01_Detail_Quantitatif"]
        self.assertEqual(ws1.cell(9, 7).number_format, "#,##0.00")


class TestFeuille1(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        import tempfile
        cls._td = tempfile.TemporaryDirectory()
        cls.addClassCleanup(cls._td.cleanup)
        from generators.modele_metre import generer_modele
        out = os.path.join(cls._td.name, "m.xlsx")
        generer_modele(out)
        cls.wb = openpyxl.load_workbook(out, data_only=False)
        cls.ws = cls.wb["01_Detail_Quantitatif"]

    def test_trois_blocs_articles(self):
        nums = [self.ws.cell(r, 1).value
                for r in (7, 19, 31)]
        self.assertEqual(nums, [1, 2, 3])

    def test_formule_qte_partielle(self):
        f = self.ws.cell(9, 10).value
        self.assertEqual(
            f, "=IF(COUNTA(G9:I9)>0,F9*PRODUCT(G9:I9),F9)")

    def test_sous_totaux_somment_j(self):
        # K du sous-total doit sommer les PARTIELS J (pas les K vides)
        self.assertEqual(self.ws.cell(17, 11).value, "=SUM(J9:J16)")
        self.assertEqual(self.ws.cell(29, 11).value, "=SUM(J21:J28)")
        self.assertEqual(self.ws.cell(41, 11).value, "=SUM(J33:J40)")

    def test_cartouche_champs_presents(self):
        textes = " ".join(
            str(self.ws.cell(r, c).value or "")
            for r in (2, 3) for c in range(1, 12))
        for mot in ("Marché", "Projet", "Bâtiment", "Date",
                    "Établi", "Vérifié"):
            self.assertIn(mot, textes)


class TestFeuille2(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        import tempfile
        cls._td = tempfile.TemporaryDirectory()
        cls.addClassCleanup(cls._td.cleanup)
        from generators.modele_metre import generer_modele
        out = os.path.join(cls._td.name, "m.xlsx")
        generer_modele(out)
        cls.wb = openpyxl.load_workbook(out, data_only=False)
        cls.ws = cls.wb["02_Armatures"]

    def test_en_tetes_diametres(self):
        for i, d in enumerate(DIAMS):
            self.assertEqual(self.ws.cell(3, 11 + i).value, f"T{d}")

    def test_ventilation_if_majuscules(self):
        f = self.ws.cell(4, 11).value
        self.assertEqual(f, '=IF($I4=6,$G4*$H4*$J4,"")')
        f = self.ws.cell(4, 19).value
        self.assertEqual(f, '=IF($I4=32,$G4*$H4*$J4,"")')

    def test_recap_longueur_poids_total_tonnes(self):
        vals = [self.ws.cell(r, 1).value for r in range(1, 30)]
        i_tot = vals.index("LONGUEUR TOTALE (ml)")
        self.assertEqual(self.ws.cell(i_tot + 1, 11).value,
                         f"=SUM(K4:K{3 + 16})")
        # Poids nominaux exacts (constantes BET)
        for i, p in enumerate(POIDS_NOMINAUX):
            self.assertAlmostEqual(self.ws.cell(i_tot + 2, 11 + i).value, p)
        i_kg = vals.index("POIDS TOTAL DES ACIERS (KG)")
        fk = self.ws.cell(i_kg + 1, 11).value
        self.assertTrue(fk.startswith("=SUM("))
        i_t = vals.index("POIDS TOTAL (tonnes)")
        self.assertIn("/1000", self.ws.cell(i_t + 1, 11).value)

    def test_taux_ferraillage_sans_div_zero(self):
        i_taux = [self.ws.cell(r, 1).value
                  for r in range(1, 30)].index(
                      "TAUX DE FERRAILLAGE MOYEN (kg/m³)")
        f = self.ws.cell(i_taux + 1, 11).value
        self.assertIn("IF(", f)


class TestFeuille3(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        import tempfile
        cls._td = tempfile.TemporaryDirectory()
        cls.addClassCleanup(cls._td.cleanup)
        from generators.modele_metre import generer_modele
        out = os.path.join(cls._td.name, "m.xlsx")
        generer_modele(out)
        cls.wb = openpyxl.load_workbook(out, data_only=False)
        cls.ws = cls.wb["03_Attachement_Ferraillage"]

    def test_en_tetes_et_rapprochement(self):
        self.assertEqual(self.ws.cell(3, 1).value, "Type d'Ouvrage")
        self.assertEqual(self.ws.cell(3, 6).value, "Longueur unitaire (m)")
        self.assertEqual(self.ws.cell(3, 7).value, "T6")
        self.assertEqual(self.ws.cell(3, 15).value, "T32")
        self.assertEqual(self.ws.cell(4, 7).value,
                         '=IF($E4=6,$C4*$D4*$F4,"")')

    def test_poids_synthese_avec_nominaux(self):
        f = self.ws.cell(18, 7).value
        self.assertTrue(f.startswith("="))
        self.assertIn("*0.222", f)


class TestFeuille4(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        import tempfile
        cls._td = tempfile.TemporaryDirectory()
        cls.addClassCleanup(cls._td.cleanup)
        from generators.modele_metre import generer_modele
        out = os.path.join(cls._td.name, "m.xlsx")
        generer_modele(out)
        cls.wb = openpyxl.load_workbook(out, data_only=False)
        cls.ws = cls.wb["04_GO_Attachement"]

    def test_quatre_prix_types(self):
        self.assertEqual(self.ws.cell(4, 2).value,
                         "Terrassement en fouilles")
        self.assertEqual(self.ws.cell(7, 2).value,
                         "Aciers Haute Adhérence FeE500")

    def test_qte_marche_vierge(self):
        """Aucune donnee projet en dur : Qte Marche vide."""
        for r in (4, 5, 6, 7):
            self.assertIsNone(self.ws.cell(r, 4).value)

    def test_liens_dynamiques(self):
        e4 = self.ws.cell(4, 5).value
        self.assertIn("01_Detail_Quantitatif", e4)
        e7 = self.ws.cell(7, 5).value
        self.assertIn("02_Armatures", e7)

    def test_ecart_et_pourcentage(self):
        self.assertEqual(self.ws.cell(4, 6).value, "=D4-E4")
        self.assertEqual(self.ws.cell(4, 7).value, "=IF(D4>0,E4/D4,0)")
        self.assertEqual(self.ws.cell(4, 7).number_format, "0.0%")


class TestVBA(unittest.TestCase):
    def test_module_bas_existe_et_complet(self):
        self.assertTrue(os.path.exists(VBA_PATH), "vba/Module1.bas manquant")
        with open(VBA_PATH, encoding="utf-8") as f:
            src = f.read()
        self.assertIn("Attribute VB_Name", src)
        for sub in ("Sub NaviguerVers(", "Sub VerifierCoherenceMetre()",
                    "Sub ReinitialiserDonneesTemplate()"):
            self.assertIn(sub, src)
        # Controles cles de l'audit
        self.assertIn("Écart acier", src)
        self.assertIn("ClearContents", src)
        # Plages alignees sur le generateur
        self.assertIn("B9:I16", src)
        self.assertIn("A4:I19", src)


if __name__ == "__main__":
    unittest.main()
