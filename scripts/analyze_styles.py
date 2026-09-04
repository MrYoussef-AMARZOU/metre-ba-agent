import openpyxl
from openpyxl.utils import get_column_letter
import sys
sys.stdout.reconfigure(encoding="utf-8")

wb = openpyxl.load_workbook("reference/metre_MZINDA.xlsx")
ws = wb["Detail quontitafif fondation"]

# Analyser les styles des cellules clés
print("=== STYLES DES CELLULES CLÉS ===")
for r in [1, 6, 7, 8, 9, 10, 11, 22, 23, 24, 33, 34, 35, 76, 100, 121, 122, 147, 148, 149, 170, 190, 215, 216, 218, 219, 245, 246, 278, 279, 317, 318, 500, 501]:
    for c in range(1, 14):
        cell = ws.cell(row=r, column=c)
        if cell.value is not None:
            font = cell.font
            fill = cell.fill
            align = cell.alignment
            border = cell.border
            print(f"  {cell.coordinate}: val={str(cell.value)[:40]!r} font={font.name}/{font.size}/{font.bold}/{font.color} fill={fill.fgColor} align={align.horizontal}/{align.wrap_text} border={border.left.style}/{border.right.style}/{border.top.style}/{border.bottom.style}")

# Largeurs de colonnes
print("\n=== LARGEURS DE COLONNES ===")
for col_letter in [get_column_letter(i) for i in range(1, 14)]:
    dim = ws.column_dimensions.get(col_letter)
    if dim:
        print(f"  {col_letter}: width={dim.width}")

# Fusions
print("\n=== CELLULES FUSIONNÉES (échantillon) ===")
for i, m in enumerate(ws.merged_cells.ranges):
    if i < 30:
        print(f"  {m}")
